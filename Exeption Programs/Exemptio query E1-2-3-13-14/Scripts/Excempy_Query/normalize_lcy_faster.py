"""
Lifecycle (LCY) file normalizer — processes .csv, .xlsx, .xls from a folder.

This faster variant folds workbook conversion and normalization into one pass:
- .xlsx/.xls are converted directly to normalized .csv
- existing .csv files still go through CSV normalization directly
- final merge behavior stays the same
"""

import csv
import json
import os
import sys
import tempfile
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, time as dt_time, timedelta
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Set, Tuple

from openpyxl import load_workbook

from merge_files import merge_files_in_folder
from lc_normalization_config import get_normalization_groups
from header_matching import normalize_header_match
from lane_labels import canonicalize_lane_label


def _require_win32():
    """Lazy-load pywin32 only when .xls COM conversion is needed."""
    try:
        import pythoncom
        from win32com.client import DispatchEx
    except ImportError as exc:
        raise ImportError(
            "pywin32 is required to process .xls files. "
            "Install it in this environment with: pip install pywin32"
        ) from exc
    return pythoncom, DispatchEx


BASE_DIR = Path(__file__).resolve().parent
_DEFAULT_INPUT_FOLDER = BASE_DIR / "Daroda_lc_vrn_files/etc"
_DEFAULT_MERGE_OUTPUT_FILE = BASE_DIR / "normalized_and_merged_etc_daroda.csv"

INPUT_FOLDER = Path(
    os.environ.get("MERGE_NORMALIZE_INPUT_FOLDER", str(_DEFAULT_INPUT_FOLDER))
).resolve()
MERGE_OUTPUT_FILE = Path(
    os.environ.get("MERGE_NORMALIZE_OUTPUT_FILE", str(_DEFAULT_MERGE_OUTPUT_FILE))
).resolve()

MERGE_AFTER_NORMALIZATION = True
MAX_WORKER_THREADS = 5
SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xls"}
ENABLE_PHASE_2 = True
SETTLED_VALUE = "SETTLED"
ACCEPTED_VALUE = "Accepted"
MIN_MERGE_HEADER_KEYWORDS = 3
OUTPUT_HEADERS = ("Date & Time", "Veh Reg No.", "MOP", "Lane No")
MERGE_HEADER_KEYWORDS = list(OUTPUT_HEADERS)
XLSX_STREAM_MAX_ROWS = 1_048_576
XLSX_STREAM_TAIL_EMPTY_ROWS = 200


def normalize_text(value) -> str:
    """Normalize header/alias text the same way as site-wide header keyword matching."""
    return normalize_header_match(value)


def build_header_lookup(groups: Dict[str, List[str]]) -> Dict[str, str]:
    lookup = {}
    for canonical, aliases in groups.items():
        lookup[normalize_text(canonical)] = canonical
        for alias in aliases:
            lookup[normalize_text(alias)] = canonical
    return lookup


def resolve_header_cell(value, lookup: Dict[str, str]) -> Optional[str]:
    return lookup.get(normalize_text(value))


def build_column_map_from_header_row(
    header_row: List, lookup: Dict[str, str]
     ) -> Dict[int, str]:
    col_map: Dict[int, str] = {}
    for col_idx, cell in enumerate(header_row, start=1):
        resolved = resolve_header_cell(cell, lookup)
        if resolved:
            col_map[col_idx] = resolved
    return col_map


def _excel_serial_to_datetime(serial: float) -> datetime:
    base = datetime(1899, 12, 30)
    return base + timedelta(days=float(serial))


def parse_datetime_cell(value) -> Optional[datetime]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date) and not isinstance(value, datetime):
        return datetime.combine(value, dt_time.min)
    if isinstance(value, (int, float)):
        try:
            return _excel_serial_to_datetime(float(value))
        except (ValueError, OSError, OverflowError):
            return None
    text = str(value).strip()
    if not text:
        return None
    fmts = (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%d-%m-%Y %H:%M:%S",
        "%d-%m-%Y %H:%M",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%m/%d/%Y %H:%M:%S",
        "%m/%d/%Y %H:%M",
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%m/%d/%Y",
    )
    for fmt in fmts:
        try:
            return datetime.strptime(text[:26], fmt)
        except ValueError:
            continue
    return None


def coerce_datetime(value) -> Optional[datetime]:
    if isinstance(value, datetime):
        return None
    return parse_datetime_cell(value)


def col_idx(col_map: Dict[int, str], name: str) -> Optional[int]:
    return next((i for i, n in col_map.items() if n == name), None)


def cell_at(row: List, idx: Optional[int]):
    if idx is None or idx > len(row):
        return None
    return row[idx - 1]


def is_settled_row(row: List, col_map: Dict[int, str]) -> bool:
    st_idx = col_idx(col_map, "Settlement Type")
    if st_idx is not None:
        return normalize_text(cell_at(row, st_idx)) == normalize_text(SETTLED_VALUE)

    # Fallback: some files don't have "Settlement Type" but provide "Transaction Status".
    ts_idx = col_idx(col_map, "Transaction Status")
    if ts_idx is not None:
        return normalize_text(cell_at(row, ts_idx)) == normalize_text(ACCEPTED_VALUE)

    # Already-normalized LC/ETC exports (Date & Time / Veh Reg No. / MOP / Lane No)
    # have no status column. Keep rows; do not treat missing status as "reject all".
    return True


def lane_as_text(value) -> str:
    return canonicalize_lane_label(value)


def veh_as_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def build_output_row(row: List, col_map: Dict[int, str]) -> Optional[Tuple]:
    if not is_settled_row(row, col_map):
        return None

    di = col_idx(col_map, "Date & Time")
    vi = col_idx(col_map, "Veh Reg No.")
    li = col_idx(col_map, "Lane No")

    vrn = veh_as_text(cell_at(row, vi))
    if not vrn:
        return None

    raw_date = cell_at(row, di)
    if isinstance(raw_date, datetime):
        dt_out: object = raw_date
    else:
        dt_out = coerce_datetime(raw_date) or parse_datetime_cell(raw_date)
        if dt_out is None:
            dt_out = raw_date

    lane = lane_as_text(cell_at(row, li))

    return (dt_out, vrn, "ETC", lane)


def get_last_used_row_col(ws):
    xl_formulas = -4123
    xl_by_rows = 1
    xl_by_columns = 2
    xl_previous = 2
    last_row_cell = ws.Cells.Find(
        What="*", LookIn=xl_formulas, SearchOrder=xl_by_rows, SearchDirection=xl_previous
    )
    last_col_cell = ws.Cells.Find(
        What="*", LookIn=xl_formulas, SearchOrder=xl_by_columns, SearchDirection=xl_previous
    )
    if last_row_cell is None or last_col_cell is None:
        return 0, 0
    return last_row_cell.Row, last_col_cell.Column


def normalize_csv_file(file_path: Path, lookup: Dict[str, str]) -> int:
    encoding = _detect_encoding(file_path)
    data_rows_written = 0
    temp_fd, temp_path = tempfile.mkstemp(
        suffix=".csv",
        prefix=f".{file_path.stem}_",
        dir=str(file_path.parent),
    )
    os.close(temp_fd)

    with file_path.open("r", encoding=encoding, newline="") as src, open(
        temp_path, "w", encoding=encoding, newline=""
    ) as dst:
        reader = csv.reader(src)
        writer = csv.writer(dst)
        header_row = next(reader, None)
        if not header_row:
            os.remove(temp_path)
            return 0

        col_map = build_column_map_from_header_row(header_row, lookup)
        if (
            col_idx(col_map, "Settlement Type") is None
            and col_idx(col_map, "Transaction Status") is None
        ):
            print(
                f"[INFO] {file_path.name}: no 'Settlement Type' or 'Transaction Status' "
                "column; keeping rows with a vehicle number (pre-normalized input)."
            )

        writer.writerow(OUTPUT_HEADERS)

        source_data_rows = 0
        for row in reader:
            source_data_rows += 1
            out = build_output_row(list(row), col_map)
            if out is None:
                continue
            date_v, vrn_v, mop_v, lane_v = out
            if isinstance(date_v, datetime):
                date_write = date_v.isoformat(sep=" ")
            else:
                date_write = date_v if date_v is not None else ""
            writer.writerow([date_write, vrn_v, mop_v, lane_v])
            data_rows_written += 1

    if source_data_rows > 0 and data_rows_written == 0:
        try:
            os.remove(temp_path)
        except OSError:
            pass
        raise RuntimeError(
            f"{file_path.name}: {source_data_rows} source row(s) produced 0 normalized "
            "rows (check Settlement Type/Transaction Status values and VRN). "
            "Source file left unchanged."
        )

    os.replace(temp_path, file_path)
    return data_rows_written


def normalize_xlsx_file(file_path: Path, lookup: Dict[str, str]) -> int:
    print(f"[START][XLSX] {file_path.name}")
    wb = load_workbook(file_path, keep_links=False)
    total_changes = 0
    try:
        for ws in wb.worksheets:
            col_map: Dict[int, str] = {}
            data_rows: List[Tuple] = []

            for r_idx, row in enumerate(
                ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column),
                start=1,
            ):
                vals = [c.value for c in row]
                if r_idx == 1:
                    col_map = build_column_map_from_header_row(
                        [("" if v is None else v) for v in vals], lookup
                    )
                    continue
                out = build_output_row(vals, col_map)
                if out is not None:
                    data_rows.append(out)

            if (
                col_idx(col_map, "Settlement Type") is None
                and col_idx(col_map, "Transaction Status") is None
            ):
                print(
                    f"[INFO] {file_path.name} -> {ws.title}: "
                    "no 'Settlement Type' or 'Transaction Status'; "
                    "keeping rows with a vehicle number (pre-normalized input)."
                )

            if ws.max_row > 1:
                ws.delete_rows(2, ws.max_row - 1)
                total_changes += 1

            while ws.max_column > 4:
                ws.delete_cols(ws.max_column, 1)
                total_changes += 1

            for c, h in enumerate(OUTPUT_HEADERS, start=1):
                ws.cell(row=1, column=c).value = h
                total_changes += 1

            for i, tup in enumerate(data_rows, start=2):
                d, v, m, lane = tup
                ws.cell(row=i, column=1).value = d
                ws.cell(row=i, column=2).value = v
                ws.cell(row=i, column=3).value = m
                ws.cell(row=i, column=4).value = str(lane)
                total_changes += 4

            print(
                f"[OK] {file_path.name} -> {ws.title}: "
                f"{len(data_rows)} data row(s), columns {list(OUTPUT_HEADERS)}"
            )

        wb.save(file_path)
        print(f"[SAVE][XLSX] {file_path.name}")
    finally:
        wb.close()
    return total_changes


def normalize_xls_file(file_path: Path, lookup: Dict[str, str]) -> int:
    pythoncom, DispatchEx = _require_win32()
    pythoncom.CoInitialize()
    excel = None
    workbook = None
    total_changes = 0
    try:
        print(f"[START][XLS] {file_path.name}")
        excel = DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        try:
            excel.ScreenUpdating = False
        except Exception:
            pass
        workbook = excel.Workbooks.Open(str(file_path))

        for ws in workbook.Worksheets:
            last_row, last_col = get_last_used_row_col(ws)
            if last_row == 0 or last_col == 0:
                continue

            col_map: Dict[int, str] = {}
            data_rows: List[Tuple] = []

            for r_idx in range(1, last_row + 1):
                row_vals = [ws.Cells(r_idx, c).Text for c in range(1, last_col + 1)]
                if r_idx == 1:
                    col_map = build_column_map_from_header_row(row_vals, lookup)
                    continue
                out = build_output_row(row_vals, col_map)
                if out is not None:
                    d, v, m, lane = out
                    if isinstance(d, datetime):
                        d_write = d
                    else:
                        d_write = parse_datetime_cell(d) if d is not None else d
                    data_rows.append((d_write, v, m, str(lane)))

            if (
                col_idx(col_map, "Settlement Type") is None
                and col_idx(col_map, "Transaction Status") is None
            ):
                print(
                    f"[INFO] {file_path.name} -> {ws.Name}: "
                    "no 'Settlement Type' or 'Transaction Status'; "
                    "keeping rows with a vehicle number (pre-normalized input)."
                )

            if last_row > 1:
                ws.Rows(f"2:{last_row}").Delete()
                total_changes += 1

            for c in range(last_col, 4, -1):
                ws.Columns(c).Delete()
                total_changes += 1

            for c, h in enumerate(OUTPUT_HEADERS, start=1):
                ws.Cells(1, c).Value = h
                total_changes += 1

            for i, tup in enumerate(data_rows, start=2):
                ws.Cells(i, 1).Value = tup[0]
                ws.Cells(i, 2).Value = tup[1]
                ws.Cells(i, 3).Value = tup[2]
                ws.Cells(i, 4).Value = tup[3]
                total_changes += 4

            print(
                f"[OK] {file_path.name} -> {ws.Name}: "
                f"{len(data_rows)} data row(s), columns {list(OUTPUT_HEADERS)}"
            )

        workbook.Save()
        workbook.Close(SaveChanges=False)
        excel.Quit()
        print(f"[SAVE][XLS] {file_path.name}")
    finally:
        try:
            if workbook is not None:
                workbook.Close(SaveChanges=False)
        except Exception:
            pass
        try:
            if excel is not None:
                excel.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()
    return total_changes


def _detect_encoding(file_path: Path) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            with file_path.open("r", encoding=enc, newline="") as f:
                f.read(4096)
            return enc
        except UnicodeDecodeError:
            continue
    return "latin-1"


def find_input_files(folder: Path):
    return sorted(
        p
        for p in folder.iterdir()
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    )


def _delete_with_retries(path: Path, retries: int = 5, delay_seconds: float = 0.8) -> None:
    last_exc: Optional[Exception] = None
    for attempt in range(1, retries + 1):
        try:
            path.unlink()
            return
        except FileNotFoundError:
            return
        except PermissionError as exc:
            last_exc = exc
            if attempt < retries:
                print(
                    f"[DEL-RETRY] {path.name}: locked (attempt {attempt}/{retries}), "
                    f"retrying in {delay_seconds:.1f}s..."
                )
                time.sleep(delay_seconds)
                continue
            raise
        except OSError as exc:
            last_exc = exc
            if attempt < retries:
                print(
                    f"[DEL-RETRY] {path.name}: {exc} (attempt {attempt}/{retries}), "
                    f"retrying in {delay_seconds:.1f}s..."
                )
                time.sleep(delay_seconds)
                continue
            raise
    if last_exc is not None:
        raise last_exc


def _normalized_csv_row_from_source(
    row: List,
    col_map: Dict[int, str],
     ) -> Optional[List[object]]:
    out = build_output_row(row, col_map)
    if out is None:
        return None

    date_v, vrn_v, mop_v, lane_v = out
    if isinstance(date_v, datetime):
        date_write: object = date_v.isoformat(sep=" ")
    else:
        date_write = date_v if date_v is not None else ""
    return [date_write, vrn_v, mop_v, str(lane_v)]


def _append_normalized_xlsx_rows(
    writer,
    ws,
    first_row_inclusive: int,
    out_cols: int,
    lookup: Dict[str, str],
    col_map: Dict[int, str],
    sheet_label: str = "",
     ) -> int:
    rows_written = 0
    tail_empty = 0
    scan_last = min(first_row_inclusive + XLSX_STREAM_MAX_ROWS - 1, 1_048_576)

    for r_idx, row in enumerate(
        ws.iter_rows(
            min_row=1,
            max_row=scan_last,
            min_col=1,
            max_col=out_cols,
            values_only=True,
        ),
        start=1,
    ):
        if r_idx < first_row_inclusive:
            continue

        row_values = list(row)[:out_cols]
        while len(row_values) < out_cols:
            row_values.append(None)

        has_data = any(normalize_text(c) != "" for c in row_values)
        if not has_data:
            tail_empty += 1
            if tail_empty >= XLSX_STREAM_TAIL_EMPTY_ROWS:
                break
        else:
            tail_empty = 0

        normalized = _normalized_csv_row_from_source(row_values, col_map)
        if normalized is None:
            continue

        writer.writerow(normalized)
        rows_written += 1

    if sheet_label:
        print(
            f"[XLSX][WRITE] {sheet_label}: wrote {rows_written} normalized row(s) "
            f"from row {first_row_inclusive} onward",
            flush=True,
        )

    return rows_written


def convert_xlsx_to_csv_with_normalization(
    path: Path,
    out_path: Path,
    keywords: Sequence[str],
    lookup: Dict[str, str],
    min_keyword_matches: Optional[int] = None,
     ) -> int:
    from csv_converter import HEADER_SCAN_MAX_ROW, detect_header_row_on_worksheet

    path = path.resolve()
    out_path = Path(out_path).resolve()
    min_matches = (
        min_keyword_matches
        if min_keyword_matches is not None
        else MIN_MERGE_HEADER_KEYWORDS
    )

    from csv_converter import WorkbookHeaderNotFoundError

    sheet_specs: List[Tuple[int, int, int]] = []
    skipped_sheets: List[str] = []
    wb1 = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        for si, ws in enumerate(wb1.worksheets):
            sheet_name = getattr(ws, "title", f"Sheet{si}")
            print(
                f"[XLSX][SCAN] sheet {si} ({sheet_name}): "
                f"checking rows 1-{HEADER_SCAN_MAX_ROW} for header..."
            )
            det = detect_header_row_on_worksheet(ws, keywords, min_matches)
            if det is None:
                skipped_sheets.append(sheet_name)
                print(
                    f"[XLSX][SKIP] sheet {si} ({sheet_name}): no row in 1-"
                    f"{HEADER_SCAN_MAX_ROW} matched at least {min_matches} "
                    "header keyword(s); skipping sheet.",
                    flush=True,
                )
                continue
            header_row, header_cols, score = det
            print(
                f"[XLSX][HEADER] sheet {si} ({sheet_name}): "
                f"row {header_row}, score={score}, cols={header_cols}"
            )
            sheet_specs.append((si, header_row, header_cols))
    finally:
        wb1.close()

    if not sheet_specs:
        skipped = ", ".join(repr(s) for s in skipped_sheets) or "(none)"
        raise WorkbookHeaderNotFoundError(
            f"{path.name}: no sheet matched at least {min_matches} header "
            f"keyword(s) in rows 1-{HEADER_SCAN_MAX_ROW}. "
            f"Skipped sheet(s): {skipped}. "
            "Adjust header keywords or sheet layout."
        )

    out_cols = max(p[2] for p in sheet_specs)
    rows_written = 0
    wb2 = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        with out_path.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(OUTPUT_HEADERS)

            first_sheet = True
            for si, header_row, _header_cols in sheet_specs:
                ws = wb2.worksheets[si]
                sheet_name = getattr(ws, "title", f"Sheet{si}")
                start = header_row + 1 if first_sheet else header_row + 1
                first_sheet = False

                header_row_values = []
                for row in ws.iter_rows(
                    min_row=header_row,
                    max_row=header_row,
                    min_col=1,
                    max_col=out_cols,
                    values_only=True,
                ):
                    header_row_values = list(row)[:out_cols]
                    break
                while len(header_row_values) < out_cols:
                    header_row_values.append(None)
                col_map = build_column_map_from_header_row(
                    [("" if v is None else v) for v in header_row_values],
                    lookup,
                )

                print(
                    f"[XLSX][START] sheet {si} ({sheet_name}): "
                    f"data-only, starting row {start}"
                )
                rows_written += _append_normalized_xlsx_rows(
                    writer,
                    ws,
                    start,
                    out_cols,
                    lookup,
                    col_map,
                    f"sheet {si} ({sheet_name})",
                )
    finally:
        wb2.close()

    print(
        f"[XLSX][DONE] {path.name}: total normalized rows written={rows_written}",
        flush=True,
    )
    return rows_written


def convert_xls_to_csv_with_normalization(
    path: Path,
    out_path: Path,
    keywords: Sequence[str],
    lookup: Dict[str, str],
    min_keyword_matches: Optional[int] = None,
      ) -> int:
    from csv_converter import detect_header_xls

    header_row, header_cols = detect_header_xls(path, keywords, min_keyword_matches)

    pythoncom, DispatchEx = _require_win32()
    pythoncom.CoInitialize()
    excel = None
    workbook = None
    rows_written = 0
    try:
        excel = DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        workbook = excel.Workbooks.Open(str(path))
        ws = workbook.Worksheets(1)
        last_row, last_col = get_last_used_row_col(ws)
        end_row = last_row if last_row else header_row

        header_values = [
            ws.Cells(header_row, c_idx).Text
            for c_idx in range(1, max(header_cols, last_col, 1) + 1)
        ]
        header_values = header_values[:header_cols]
        while len(header_values) < header_cols:
            header_values.append("")
        col_map = build_column_map_from_header_row(header_values, lookup)

        with out_path.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(OUTPUT_HEADERS)
            for r_idx in range(header_row + 1, end_row + 1):
                raw = [
                    ws.Cells(r_idx, c_idx).Text
                    for c_idx in range(1, max(header_cols, last_col, 1) + 1)
                ]
                raw = raw[:header_cols]
                while len(raw) < header_cols:
                    raw.append("")

                normalized = _normalized_csv_row_from_source(raw, col_map)
                if normalized is None:
                    continue

                writer.writerow(normalized)
                rows_written += 1
    finally:
        try:
            if workbook is not None:
                workbook.Close(SaveChanges=False)
        except Exception:
            pass
        try:
            if excel is not None:
                excel.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()

    return rows_written


def _phase1_header_keywords() -> Tuple[List[str], int]:
    raw = os.environ.get("MERGE_NORMALIZE_HEADER_KEYWORDS", "").strip()
    if raw:
        try:
            parsed = json.loads(raw)
            if isinstance(parsed, list):
                keywords = [str(k).strip() for k in parsed if str(k).strip()]
                if len(keywords) >= MIN_MERGE_HEADER_KEYWORDS:
                    return keywords, min(MIN_MERGE_HEADER_KEYWORDS, len(keywords))
        except json.JSONDecodeError:
            pass

    from csv_converter import HEADER_KEYWORDS, MIN_HEADER_KEYWORD_MATCHES

    return list(HEADER_KEYWORDS), MIN_HEADER_KEYWORD_MATCHES


def convert_workbook_to_normalized_csv(
    path: Path,
    out_path: Path,
    keywords: Sequence[str],
    lookup: Dict[str, str],
    min_keyword_matches: Optional[int] = None,
    ) -> int:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return convert_xlsx_to_csv_with_normalization(
            path,
            out_path,
            keywords,
            lookup,
            min_keyword_matches,
        )
    if suffix == ".xls":
        return convert_xls_to_csv_with_normalization(
            path,
            out_path,
            keywords,
            lookup,
            min_keyword_matches,
        )
    raise ValueError(f"Expected .xlsx or .xls, got {suffix}")


def phase_convert_workbooks_to_csv_and_delete(
    folder_path: Path,
    lookup: Dict[str, str],
    ) -> Set[Path]:
    converted_csvs: Set[Path] = set()

    try:
        from csv_converter import HEADER_SCAN_MAX_ROW, WorkbookHeaderNotFoundError
    except ImportError as exc:
        print(
            f"[WARN] Could not import csv_converter ({exc}); "
            "skipping pre-conversion to CSV.",
            flush=True,
        )
        return converted_csvs

    keywords, min_matches = _phase1_header_keywords()
    if len(keywords) < min_matches:
        print(
            f"[WARN] Need at least {min_matches} header keywords for pre-conversion; "
            f"have {len(keywords)}. Skipping pre-conversion to CSV.",
            flush=True,
        )
        return converted_csvs

    folder_path = folder_path.resolve()
    if not folder_path.is_dir():
        return converted_csvs

    targets = [
        p
        for p in sorted(folder_path.iterdir(), key=lambda x: x.name.lower())
        if p.is_file()
        and p.suffix.lower() in (".xlsx", ".xls")
        and not p.name.startswith("~$")
    ]
    if not targets:
        print("[PHASE 1] No .xlsx/.xls files to pre-convert.", flush=True)
        return converted_csvs

    print(
        f"\n[PHASE 1] Workbook(s) -> normalized .csv "
        f"(>={min_matches} keyword hits in rows 1-{HEADER_SCAN_MAX_ROW} per sheet), "
        f"{len(targets)} workbook(s)...",
        flush=True,
    )
    for path in targets:
        out_csv = path.with_suffix(".csv")
        try:
            n_rows = convert_workbook_to_normalized_csv(
                path,
                out_csv,
                keywords,
                lookup,
                min_keyword_matches=min_matches,
            )
            _delete_with_retries(path)
            converted_csvs.add(out_csv.resolve())
            print(
                f"[CONVERT+DEL] {path.name} -> {out_csv.name} "
                f"(normalized data rows written: {n_rows})",
                flush=True,
            )
        except WorkbookHeaderNotFoundError as exc:
            print(f"[CONVERT-FAIL] {path.name}: {exc}", flush=True)
            raise SystemExit(f"Header detection failed for {path.name}: {exc}") from exc
        except Exception as exc:
            print(
                f"[CONVERT-SKIP] {path.name}: {exc} "
                f"(will use default .xlsx/.xls normalization)",
                flush=True,
            )

    return converted_csvs


def process_file(file_path: Path, lookup: Dict[str, str]):
    suf = file_path.suffix.lower()
    if suf == ".csv":
        n = normalize_csv_file(file_path, lookup)
    elif suf == ".xlsx":
        n = normalize_xlsx_file(file_path, lookup)
    elif suf == ".xls":
        n = normalize_xls_file(file_path, lookup)
    else:
        return None
    return file_path.name, n


def merge_normalized_files() -> bool:
    print(f"\n[MERGE] Starting merge from {INPUT_FOLDER}")
    MERGE_OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    if MERGE_OUTPUT_FILE.exists():
        try:
            MERGE_OUTPUT_FILE.unlink()
        except OSError:
            pass
    merge_files_in_folder(
        str(INPUT_FOLDER),
        str(MERGE_OUTPUT_FILE),
        MERGE_HEADER_KEYWORDS,
    )
    if not MERGE_OUTPUT_FILE.exists():
        print(
            f"[ERROR] Merge did not create {MERGE_OUTPUT_FILE.name}. "
            "No processable LC/ETC rows found after normalization.",
            flush=True,
        )
        return False
    return True


def main():
    start = time.perf_counter()
    groups = get_normalization_groups()
    lookup = build_header_lookup(groups)
    INPUT_FOLDER.mkdir(parents=True, exist_ok=True)
    MERGE_OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    print(f"Input folder: {INPUT_FOLDER}", flush=True)
    print(f"Merge output: {MERGE_OUTPUT_FILE}", flush=True)
    print(f"Loaded {len(groups)} LC header group(s) from config.", flush=True)

    converted_csvs = phase_convert_workbooks_to_csv_and_delete(INPUT_FOLDER, lookup)

    if not ENABLE_PHASE_2:
        elapsed = time.perf_counter() - start
        print("\n[PHASE 2] Disabled (temporary switch ENABLE_PHASE_2=False).")
        print(f"Execution time: {elapsed:.2f} seconds")
        return

    files = [p for p in find_input_files(INPUT_FOLDER) if p.resolve() not in converted_csvs]
    print(f"Input folder: {INPUT_FOLDER}")
    print(f"\n[PHASE 2] Found {len(files)} file(s) to normalize")

    normalize_errors = 0
    if not files and not converted_csvs:
        elapsed = time.perf_counter() - start
        print("[INFO] Nothing to process.")
        if MERGE_AFTER_NORMALIZATION:
            ok = merge_normalized_files()
            print(f"Execution time: {elapsed:.2f} seconds")
            if not ok:
                sys.exit(1)
        else:
            print("\n[MERGE] Skipped because MERGE_AFTER_NORMALIZATION is False")
            print(f"Execution time: {elapsed:.2f} seconds")
        return

    if files:
        worker_count = min(MAX_WORKER_THREADS, len(files))
        print(
            f"Thread pool: up to {worker_count} file(s) processed concurrently "
            f"(max_workers={worker_count}, cap={MAX_WORKER_THREADS})"
        )

        with ThreadPoolExecutor(max_workers=worker_count) as executor:
            future_to_path = {
                executor.submit(process_file, file_path, lookup): file_path
                for file_path in files
            }
            for future in as_completed(future_to_path):
                file_path = future_to_path[future]
                try:
                    result = future.result()
                    if result is None:
                        continue
                    name, count = result
                    print(f"\nFile: {name}  (cells/rows written ~ {count})")
                except Exception as e:
                    normalize_errors += 1
                    print(f"\nError in {file_path.name}: {e}")

    if MERGE_AFTER_NORMALIZATION:
        ok = merge_normalized_files()
    else:
        print("\n[MERGE] Skipped because MERGE_AFTER_NORMALIZATION is False")
        ok = True

    elapsed = time.perf_counter() - start
    print(f"\nExecution time: {elapsed:.2f} seconds")
    if normalize_errors or not ok:
        sys.exit(1)


if __name__ == "__main__":
    main()
