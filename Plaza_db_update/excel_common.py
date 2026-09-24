"""Shared Excel/CSV reading, datetime parsing, and value normalization."""

from __future__ import annotations

import re
import shutil
import tempfile
from dataclasses import dataclass
from datetime import datetime, time
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from config.excel_config import (
    AM_PM_PATTERN,
    COLUMN_MAPPING,
    DATE_COLUMN_ALIASES,
    DATE_ONLY_FORMATS,
    DATETIME_FORMATS,
    DATETIME_FORMATS_12H,
    DATETIME_FORMATS_24H,
    ETC_DATETIME_COLUMN_ALIASES,
    ETC_NPCI_CLASS_COLUMN_ALIASES,
    ETC_SETTLEMENT_COLUMN_ALIASES,
    EXCEL_EXTENSIONS,
    HEADER_KEYWORDS,
    LANE_COLUMN_ALIASES,
    LANE_LOOKUP,
    MAX_SUPPORTED_LANES,
    MOP_COLUMN_ALIASES,
    TIME_COLUMN_ALIASES,
    TIME_ONLY_FORMATS,
    TIME_ONLY_FORMATS_12H,
    TIME_ONLY_FORMATS_24H,
    VEHICLE_CLASS_COLUMN_ALIASES,
    normalize_key,
    required_excel_columns,
)
from config.mappings import (
    load_mop_ignore_aliases,
    load_mop_mappings,
    load_vehicle_class_mappings,
)
from config.excel_config import split_mappings, build_lookup, normalize_key


class UnmappedValueError(Exception):
    """Raised when a value cannot be normalized; execution should stop."""


VEHICLE_CLASS_MAPPINGS = load_vehicle_class_mappings()
MOP_MAPPINGS = load_mop_mappings()
VEHICLE_CLASS_COLUMNS, VEHICLE_CLASS_NORMALIZATION = split_mappings(VEHICLE_CLASS_MAPPINGS)
MOP_COLUMNS, MOP_NORMALIZATION = split_mappings(MOP_MAPPINGS)
VEHICLE_CLASS_LOOKUP = build_lookup(VEHICLE_CLASS_NORMALIZATION)
MOP_LOOKUP = build_lookup(MOP_NORMALIZATION)
# Case-insensitive ignore set — these MOP labels are not counted and do not stop ETL.
MOP_IGNORE_LOOKUP = {
    normalize_key(alias) for alias in load_mop_ignore_aliases() if normalize_key(alias)
}


def is_blank(value) -> bool:
    if pd.isna(value):
        return True
    return str(value).strip() == ""


def datetime_has_am_pm(text: str) -> bool:
    return bool(AM_PM_PATTERN.search(text))


def datetime_formats_for_text(text: str) -> list[str]:
    """Use 12-hour formats when AM/PM is present; otherwise 24-hour formats."""
    if datetime_has_am_pm(text):
        return DATETIME_FORMATS_12H
    return DATETIME_FORMATS_24H


def time_formats_for_text(text: str) -> list[str]:
    if datetime_has_am_pm(text):
        return TIME_ONLY_FORMATS_12H
    return TIME_ONLY_FORMATS_24H


def non_empty_sample(values, *, limit: int | None = None) -> list:
    sample = [
        value
        for value in values
        if not pd.isna(value) and str(value).strip()
    ]
    if limit is not None:
        return sample[:limit]
    return sample


def try_parse_with_format(value, datetime_format: str):
    if pd.isna(value):
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()

    text = str(value).strip()
    if not text:
        return None

    if "%p" in datetime_format and not datetime_has_am_pm(text):
        return None
    if "%p" not in datetime_format and datetime_has_am_pm(text):
        return None

    try:
        return datetime.strptime(text, datetime_format)
    except ValueError:
        return None


def parse_datetime_with_format(value, datetime_format: str) -> datetime:
    if pd.isna(value):
        raise ValueError("Empty datetime value")
    if isinstance(value, datetime):
        return value
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()

    text = str(value).strip()
    return datetime.strptime(text, datetime_format)


def _best_format(sample: list, formats: list[str]) -> str | None:
    best_format = None
    best_count = -1
    for fmt in formats:
        parsed_count = sum(
            1 for value in sample if try_parse_with_format(value, fmt) is not None
        )
        if parsed_count > best_count:
            best_count = parsed_count
            best_format = fmt
    if best_format is None or best_count == 0:
        return None
    if best_count < len(sample):
        print(
            f"  Warning: detected format '{best_format}' parsed "
            f"{best_count}/{len(sample)} sample(s)."
        )
    return best_format


def detect_datetime_format(values) -> str:
    """
    Detect the datetime format used in a file by testing DATETIME_FORMATS
    against non-empty sample values.
    """
    sample = non_empty_sample(values)
    if not sample:
        raise ValueError("No datetime values found to detect format.")

    best_format = _best_format(sample, DATETIME_FORMATS)
    if best_format is None:
        raise ValueError("Could not detect datetime format from file values.")
    return best_format


def detect_date_only_format(values) -> str:
    sample = non_empty_sample(values)
    if not sample:
        raise ValueError("No date values found to detect format.")
    best_format = _best_format(sample, DATE_ONLY_FORMATS)
    if best_format is None:
        raise ValueError("Could not detect date-only format from file values.")
    return best_format


def detect_time_only_format(values) -> str:
    sample = non_empty_sample(values)
    if not sample:
        raise ValueError("No time values found to detect format.")
    best_format = _best_format(sample, TIME_ONLY_FORMATS)
    if best_format is None:
        raise ValueError("Could not detect time-only format from file values.")
    return best_format


def find_column(df: pd.DataFrame, mapped_name: str) -> str:
    if mapped_name in df.columns:
        return mapped_name

    normalized_target = normalize_key(mapped_name)
    for column in df.columns:
        if normalize_key(column) == normalized_target:
            return column

    available = ", ".join(str(col) for col in df.columns)
    raise KeyError(
        f"Mapped column '{mapped_name}' not found. Available columns: {available}"
    )


def find_column_by_aliases(df: pd.DataFrame, aliases: list[str]) -> str:
    last_error: Exception | None = None
    for alias in aliases:
        try:
            return find_column(df, alias)
        except KeyError as exc:
            last_error = exc
    available = ", ".join(str(col) for col in df.columns)
    raise KeyError(
        f"None of the column aliases {aliases!r} were found. "
        f"Available columns: {available}"
    ) from last_error


def find_optional_column_by_aliases(df: pd.DataFrame, aliases: list[str]) -> str | None:
    try:
        return find_column_by_aliases(df, aliases)
    except KeyError:
        return None


@dataclass(frozen=True)
class DatetimeResolution:
    """How event timestamps are stored in a VRN file."""

    date_col: str
    mode: str  # "combined" | "split"
    datetime_format: str | None = None
    date_format: str | None = None
    time_format: str | None = None
    time_col: str | None = None

    @property
    def label(self) -> str:
        if self.mode == "split":
            return (
                f"split date '{self.date_format}' + time '{self.time_format}' "
                f"(columns {self.date_col!r} + {self.time_col!r})"
            )
        return f"combined '{self.datetime_format}' (column {self.date_col!r})"


def resolve_datetime_columns(df: pd.DataFrame) -> DatetimeResolution:
    """
    Resolve datetime parsing for a file.

    Prefer a single combined datetime column. If DATE values are date-only,
    look for a TIME column and parse the pair.
    """
    date_col = find_column_by_aliases(df, DATE_COLUMN_ALIASES)
    date_values = df[date_col].tolist()

    try:
        datetime_format = detect_datetime_format(date_values)
        return DatetimeResolution(
            date_col=date_col,
            mode="combined",
            datetime_format=datetime_format,
        )
    except ValueError:
        pass

    try:
        date_format = detect_date_only_format(date_values)
    except ValueError as exc:
        raise ValueError(
            f"Could not detect datetime or date-only format in column {date_col!r}."
        ) from exc

    time_col = find_optional_column_by_aliases(df, TIME_COLUMN_ALIASES)
    if time_col is None:
        raise ValueError(
            f"Column {date_col!r} looks date-only ({date_format}), but no TIME "
            f"column was found (tried aliases {TIME_COLUMN_ALIASES!r})."
        )

    try:
        time_format = detect_time_only_format(df[time_col].tolist())
    except ValueError as exc:
        raise ValueError(
            f"Found time column {time_col!r}, but could not detect its format."
        ) from exc

    return DatetimeResolution(
        date_col=date_col,
        mode="split",
        date_format=date_format,
        time_format=time_format,
        time_col=time_col,
    )


def _parse_time_only(value, time_format: str) -> time | None:
    if pd.isna(value):
        return None
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime().time()
    if isinstance(value, time):
        return value

    text = str(value).strip()
    if not text:
        return None
    try:
        return datetime.strptime(text, time_format).time()
    except ValueError:
        return None


def combine_date_and_time(
    date_value,
    time_value,
    *,
    date_format: str,
    time_format: str,
) -> datetime | None:
    date_part = try_parse_with_format(date_value, date_format)
    if date_part is None and isinstance(date_value, (datetime, pd.Timestamp)):
        date_part = (
            date_value.to_pydatetime()
            if isinstance(date_value, pd.Timestamp)
            else date_value
        )
    if date_part is None:
        return None

    time_part = _parse_time_only(time_value, time_format)
    if time_part is None:
        return None

    return datetime.combine(date_part.date(), time_part)


def parse_event_datetime(date_value, resolution: DatetimeResolution, time_value=None):
    """Parse one event timestamp using a DatetimeResolution."""
    if resolution.mode == "combined":
        return try_parse_datetime(date_value, datetime_format=resolution.datetime_format)

    return combine_date_and_time(
        date_value,
        time_value,
        date_format=resolution.date_format or "",
        time_format=resolution.time_format or "",
    )


def safe_parse_event_datetime(date_value, resolution: DatetimeResolution, time_value=None):
    parsed = parse_event_datetime(date_value, resolution, time_value=time_value)
    return parsed if parsed is not None else pd.NaT


def parse_datetime(value, datetime_format: str | None = None) -> datetime:
    if pd.isna(value):
        raise ValueError("Empty datetime value")
    if isinstance(value, datetime):
        return value
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()

    if datetime_format:
        return parse_datetime_with_format(value, datetime_format)

    text = str(value).strip()
    for fmt in datetime_formats_for_text(text):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue

    parsed = pd.to_datetime(text, dayfirst=True, errors="coerce")
    if pd.isna(parsed):
        raise ValueError(f"Unable to parse datetime: {value!r}")
    return parsed.to_pydatetime()


def try_parse_datetime(value, datetime_format: str | None = None):
    try:
        return parse_datetime(value, datetime_format=datetime_format)
    except (ValueError, TypeError):
        return None


def hour_bucket_label(dt: datetime) -> str:
    """Return 24-hour bucket label, e.g. 2:23 AM -> '2-3', 2:23 PM -> '14-15'."""
    return f"{dt.hour}-{dt.hour + 1}"


def try_normalize_vehicle_class(value):
    """Map Excel vehicle class to canonical name (case-insensitive)."""
    key = normalize_key(value)
    if not key:
        return None
    return VEHICLE_CLASS_LOOKUP.get(key)


def try_normalize_mop(value):
    """Map Excel MOP to canonical name (case-insensitive). Ignored labels return None."""
    key = normalize_key(value)
    if not key:
        return None
    if key in MOP_IGNORE_LOOKUP:
        return None
    return MOP_LOOKUP.get(key)


def is_ignored_mop(value) -> bool:
    """True when MOP is listed under mop.json 'ignore' (case-insensitive)."""
    key = normalize_key(value)
    if not key:
        return False
    return key in MOP_IGNORE_LOOKUP


def try_normalize_lane(value):
    """
    Map raw Excel lane labels to canonical L01…L12.

    Only explicit aliases in LANE_MAPPINGS are accepted (e.g. L1 → L01).
    Unknown values return None so callers can stop execution.
    """
    key = normalize_key(value)
    if not key:
        return None
    return LANE_LOOKUP.get(key)


def extract_lane_number(value) -> int | None:
    """Best-effort numeric lane index from raw labels like L14 / Lane 14."""
    key = normalize_key(value)
    if not key:
        return None
    match = re.search(r"(\d+)", key)
    if not match:
        return None
    return int(match.group(1))


def unsupported_high_lane_counts(raw_lane_values: list) -> dict[str, int]:
    """
    Labels that appear to be lane numbers above MAX_SUPPORTED_LANES.
    Used only for a clearer stop message — not for remapping.
    """
    found: dict[str, int] = {}
    for value in raw_lane_values:
        if is_blank(value):
            continue
        if try_normalize_lane(value) is not None:
            continue
        lane_number = extract_lane_number(value)
        if lane_number is None or lane_number <= MAX_SUPPORTED_LANES:
            continue
        raw = str(value).strip()
        found[raw] = found.get(raw, 0) + 1
    return found


def lane_limit_error_message(file_name: str, lane_col: str, high_lane_counts: dict[str, int]) -> str:
    details = ", ".join(
        f"{label!r} ({count} row(s))" for label, count in sorted(high_lane_counts.items())
    )
    return (
        f"Plaza has lane(s) beyond supported L01–L{MAX_SUPPORTED_LANES:02d} "
        f"in '{file_name}' ({lane_col}): {details}. "
        f"Add mappings/columns for the new lanes in config (LANE_MAPPINGS / DB), then rerun. "
        f"ETL stopped — no fallback remapping is applied."
    )


def optional_normalize_vehicle_class(value):
    if is_blank(value):
        return None
    return try_normalize_vehicle_class(value)


def optional_normalize_lane(value):
    if is_blank(value):
        return None
    return try_normalize_lane(value)


def optional_normalize_mop(value):
    if is_blank(value):
        return None
    return try_normalize_mop(value)


def safe_parse_datetime(value, datetime_format: str):
    if is_blank(value):
        return pd.NaT
    try:
        return parse_datetime(value, datetime_format=datetime_format)
    except (ValueError, TypeError):
        return pd.NaT


def find_header_row(
    df: pd.DataFrame,
    required_fields: tuple[str, ...] | list[str] | None = None,
) -> int | None:
    """
    Find the header row by matching required Excel column names (including aliases)
    as exact cell values.

    This avoids false positives from report metadata rows such as
    'FROM DATE' / 'TO DATE' that only contain DATE as a substring.
    """
    fields = required_fields or list(COLUMN_MAPPING.keys())
    field_alias_groups: dict[str, list[str]] = {
        "datetime": DATE_COLUMN_ALIASES,
        "vehicle_class": VEHICLE_CLASS_COLUMN_ALIASES,
        "lane_no": LANE_COLUMN_ALIASES,
        "mop": MOP_COLUMN_ALIASES,
        "etc_datetime": ETC_DATETIME_COLUMN_ALIASES,
        "etc_npci": ETC_NPCI_CLASS_COLUMN_ALIASES,
        "etc_settlement": ETC_SETTLEMENT_COLUMN_ALIASES,
    }
    required_alias_groups: list[list[str]] = []
    for field in fields:
        aliases = field_alias_groups.get(field)
        if aliases:
            required_alias_groups.append(aliases)
        else:
            required_alias_groups.append([COLUMN_MAPPING[field]])

    best_row = None
    best_score = 0
    min_score = min(2, len(required_alias_groups))

    for row_index in range(len(df)):
        row_cells = {
            normalize_key(value)
            for value in df.iloc[row_index]
            if not pd.isna(value) and str(value).strip()
        }
        score = sum(
            1
            for aliases in required_alias_groups
            if any(normalize_key(alias) in row_cells for alias in aliases)
        )
        if score > best_score:
            best_score = score
            best_row = row_index

    if best_score >= min_score:
        return best_row

    return None


def remove_duplicate_header_rows(df: pd.DataFrame) -> pd.DataFrame:
    """Remove consecutive duplicate header-like rows from the raw sheet."""
    rows_to_remove: list[int] = []

    for row_index in range(len(df) - 1):
        current_row = df.iloc[row_index]
        next_row = df.iloc[row_index + 1]
        current_str = current_row.fillna("").astype(str)
        next_str = next_row.fillna("").astype(str)

        if not current_str.equals(next_str):
            continue

        has_text = any(
            pd.notna(value) and isinstance(value, str) and str(value).strip()
            for value in current_row
        )
        if has_text:
            rows_to_remove.append(row_index)

    if not rows_to_remove:
        return df

    return df.drop(index=rows_to_remove).reset_index(drop=True)


def unmerge_workbook(input_path: Path, output_path: Path, sheet_name: str | None = None) -> str:
    """Unmerge all merged cells and save to a new workbook."""
    workbook = load_workbook(input_path)
    try:
        worksheet = workbook[sheet_name] if sheet_name else workbook.active
        resolved_sheet_name = worksheet.title

        for merged_range in list(worksheet.merged_cells):
            worksheet.unmerge_cells(range_string=str(merged_range))

        workbook.save(output_path)
        return resolved_sheet_name
    finally:
        workbook.close()


def remove_empty_columns(filename: Path, sheet_name: str, header_row: int) -> Path:
    """Delete columns with empty headers or no data below the header row."""
    workbook = load_workbook(filename)
    try:
        worksheet = workbook[sheet_name]
        max_col = worksheet.max_column
        cols_to_delete: list[int] = []

        for col_idx in range(1, max_col + 1):
            col_letter = get_column_letter(col_idx)
            header_value = worksheet[f"{col_letter}{header_row}"].value

            if header_value is None or str(header_value).strip() == "":
                cols_to_delete.append(col_idx)
                continue

            empty_below = True
            for row_idx in range(header_row + 1, worksheet.max_row + 1):
                value = worksheet.cell(row=row_idx, column=col_idx).value
                if value not in (None, ""):
                    empty_below = False
                    break

            if empty_below:
                cols_to_delete.append(col_idx)

        for col_idx in sorted(cols_to_delete, reverse=True):
            worksheet.delete_cols(col_idx, 1)

        output_path = filename.with_name(f"{filename.stem}_cleaned{filename.suffix}")
        workbook.save(output_path)
        return output_path
    finally:
        workbook.close()


def read_csv_with_header_detection(path: Path) -> pd.DataFrame:
    """Read CSV after detecting and trimming rows above the real header."""
    keywords = [keyword.upper() for keyword in HEADER_KEYWORDS]

    with open(path, "r", encoding="utf-8", errors="ignore") as file_obj:
        lines = file_obj.readlines()

    detected_header_index = None
    for line_index, line in enumerate(lines):
        upper_line = line.upper()
        keyword_count = sum(1 for keyword in keywords if keyword in upper_line)
        if keyword_count >= 2:
            detected_header_index = line_index
            break

    if detected_header_index is None:
        raise ValueError(f"Header could not be detected in CSV: {path.name}")

    print(f"  Header detected at CSV line: {detected_header_index + 1}")
    trimmed_lines = lines[detected_header_index:]

    temp_clean = path.with_name(f"{path.stem}_clean.csv")
    try:
        with open(temp_clean, "w", encoding="utf-8") as file_obj:
            file_obj.writelines(trimmed_lines)

        try:
            return pd.read_csv(temp_clean, engine="python", sep=None, dtype=str)
        except Exception:
            return pd.read_csv(temp_clean, engine="python", on_bad_lines="skip", dtype=str)
    finally:
        if temp_clean.exists():
            temp_clean.unlink()


def prepare_excel_dataframe(
    file_path: Path,
    required_fields: tuple[str, ...] | list[str] | None = None,
) -> pd.DataFrame:
    """
    Normalize Excel files before column mapping:
    1) remove duplicate header rows
    2) detect header row
    3) unmerge merged cells
    4) remove empty columns
    5) read data using detected header row
    """
    with tempfile.TemporaryDirectory() as temp_dir:
        temp_dir_path = Path(temp_dir)
        working_copy = temp_dir_path / file_path.name
        shutil.copy2(file_path, working_copy)

        df_full = pd.read_excel(working_copy, header=None)
        df_full = remove_duplicate_header_rows(df_full)

        header_index = find_header_row(df_full, required_fields=required_fields)
        if header_index is None:
            raise ValueError(f"Header row could not be detected in: {file_path.name}")

        print(f"  Header detected at Excel row: {header_index + 1}")

        no_dup_path = temp_dir_path / "no_duplicate_headers.xlsx"
        df_full.to_excel(no_dup_path, index=False, header=False)

        # Close ExcelFile before the next open — Windows locks the xlsx otherwise.
        with pd.ExcelFile(no_dup_path) as excel_file:
            sheet_name = excel_file.sheet_names[0]

        unmerged_path = temp_dir_path / "unmerged.xlsx"
        sheet_name = unmerge_workbook(no_dup_path, unmerged_path, sheet_name=sheet_name)

        cleaned_path = remove_empty_columns(
            unmerged_path,
            sheet_name,
            header_index + 1,
        )

        df = pd.read_excel(
            cleaned_path,
            sheet_name=sheet_name,
            header=header_index,
            dtype=str,
        )

    df.columns = [str(column).strip() for column in df.columns]
    df = df.dropna(how="all").reset_index(drop=True)
    return df


def read_excel_file(
    file_path: Path,
    required_fields: tuple[str, ...] | list[str] | None = None,
) -> pd.DataFrame:
    suffix = file_path.suffix.lower()
    if suffix == ".csv":
        df = read_csv_with_header_detection(file_path)
    else:
        df = prepare_excel_dataframe(file_path, required_fields=required_fields)

    df.columns = [str(column).strip() for column in df.columns]
    return df.dropna(how="all").reset_index(drop=True)


def list_excel_files(folder_path: Path) -> list[Path]:
    """Return Excel/CSV files under folder_path, including all subfolders."""
    return sorted(
        path
        for path in folder_path.rglob("*")
        if path.is_file()
        and path.suffix.lower() in EXCEL_EXTENSIONS
        and not path.name.startswith("~$")
    )


def is_header_detection_error(exc: Exception) -> bool:
    message = str(exc).lower()
    return "header" in message and "detect" in message


def is_unreadable_workbook_error(exc: Exception) -> bool:
    """True when Excel/openpyxl cannot open a corrupt or invalid workbook."""
    message = str(exc).lower()
    needles = (
        "could not read stylesheet",
        "unable to read workbook",
        "invalid xml",
        "not a zip file",
        "bad zip file",
        "file is not a zip file",
        "workbook source files contain some invalid",
        "there is no item named 'xl/styles",
        "error reading existing file",
        "content_types",
        "does not support file format",
        "excel file format cannot be determined",
    )
    return any(needle in message for needle in needles)


def is_skippable_excel_read_error(exc: Exception) -> bool:
    """Header-detection failures and corrupt workbooks should not stop the batch."""
    return is_header_detection_error(exc) or is_unreadable_workbook_error(exc)


def create_run_log_file(insights_dir: Path, module_name: str) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_path = insights_dir / f"log-{module_name}-{timestamp}.txt"
    log_path.write_text(
        f"{module_name} processing log\n"
        f"Started: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
        f"{'=' * 80}\n",
        encoding="utf-8",
    )
    return log_path


def append_run_log(log_path: Path, message: str) -> None:
    with log_path.open("a", encoding="utf-8") as log_file:
        log_file.write(message + "\n")
